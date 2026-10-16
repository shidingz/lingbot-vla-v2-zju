"""Create one editable slide and matching local preview, without touching training."""
import sys, json, math
from pathlib import Path
from datetime import datetime
sys.path.insert(0, '/tmp/lingbot_pptx_deps')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parent
D=json.loads((ROOT/'metrics_snapshot.json').read_text())
T,N=D['training'],D['npu']; latest=T['latest_loss']; step=latest['step']; total=T['total_steps']
sec=T['recent_seconds_per_step']; old=T['baseline']['last500_wall_seconds_per_step']; saving=(1-sec/old)*100
stamp=datetime.fromisoformat(D['captured_at']).strftime('%m月%d日 %H:%M')
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
slide=prs.slides.add_slide(prs.slide_layouts[6])
prs.core_properties.title='LingBot-VLA 2.0｜训练策略、速度与资源观察'
prs.core_properties.subject='当前一轮训练的团队讨论报告'
prs.core_properties.author='LingBot 训练实验'
prs.core_properties.keywords='LingBot, RoboTwin, Ascend, 1 epoch, training'
W,H=1600,900;S=1.5
im=Image.new('RGB',(int(W*S),int(H*S)), '#F3F6FA');draw=ImageDraw.Draw(im)
FONT='/tmp/lingbot_report_fonts/NotoSansSC-Regular.ttf'
INK='183047'; MUTED='627386'; TEAL='087E74'; BLUE='285DB4'; LINE='E3EAF0'; BG='F3F6FA'; WHITE='FFFFFF'; AMBER='9C5B11'
objects=[]; warnings=[]
def rgb(c):return RGBColor.from_string(c)
def font(size):return ImageFont.truetype(FONT,round(size*100/72*S))
def rect(x,y,w,h,fill,radius=0,line=None):
 sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.6)
 else:sh.line.fill.background()
 if radius:
  sh.adjustments[0]=.10
  draw.rounded_rectangle((x*S,y*S,(x+w)*S,(y+h)*S),radius=radius*S,fill='#'+fill,outline='#'+line if line else None)
 else:draw.rectangle((x*S,y*S,(x+w)*S,(y+h)*S),fill='#'+fill,outline='#'+line if line else None)
 objects.append(('rect',x,y,w,h))
 return sh

def txt(x,y,w,h,text,size=14,color=INK,bold=False,align='left',leading=None):
 box=slide.shapes.add_textbox(Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100))
 tf=box.text_frame;tf.clear();tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
 lines=text.split('\n'); line_px=(leading or size*100/72*(1.15 if len(lines)==1 else 1.35))
 for j,s in enumerate(lines):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(line_px*72/100)
  p.alignment=PP_ALIGN.CENTER if align=='center' else PP_ALIGN.RIGHT if align=='right' else PP_ALIGN.LEFT
  run=p.add_run();run.text=s;run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)
  ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
  f=font(size);length=draw.textlength(s,font=f)/S
  if length>w+1:warnings.append({'text':s,'width':length,'box':w})
  xx=x+(w-length)/2 if align=='center' else x+w-length if align=='right' else x
  draw.text((xx*S,(y+j*line_px+2)*S),s,font=f,fill='#'+color,anchor='lt',stroke_width=0)
 if len(lines)*line_px>h+5:warnings.append({'height_text':text,'needed':len(lines)*line_px,'box':h})
 objects.append(('text',x,y,w,h,text));return box

rect(0,0,W,H,BG)
rect(0,0,W,112,INK)
txt(48,23,1504,50,'LingBot-VLA 2.0｜训练策略、速度与资源观察',30,WHITE,True)
txt(50,81,1500,28,f'运行 robotwin_1ep_fast  ·  RoboTwin：50 任务 × 50 轨迹（clean 来源待核验）  ·  {stamp} 北京时间',12,'C8D7E4')

cards=[('训练计划','1 epoch','保持动作专家训练与辅助目标'),('单轮计算预算',f'≈{T["estimated_epoch_hours"]:.0f} 小时','按近期 100 步实测速度外推'),('当前训练吞吐',f'{96/sec:.1f} 样本 / 秒',f'全局 batch 96 · {sec:.2f} 秒 / 步'),('资源观察','显存紧 · 利用率低','计算未持续跑满，仍有优化空间')]
for i,(label,value,note) in enumerate(cards):
 x=48+i*380
 rect(x,131,364,113,'E5F3F0' if i==1 else WHITE,12)
 txt(x+18,141,328,25,label,12,MUTED)
 txt(x+18,170,328,44,value,27 if i!=3 else 22,TEAL if i==1 else INK,True)
 txt(x+18,215,328,23,note,11.5,MUTED)

for x,w in [(48,500),(572,474),(1070,482)]:rect(x,267,w,450,WHITE,12)
for x,title in [(68,'01  训练什么'),(592,'02  速度与实验进度'),(1090,'03  八卡资源利用')]:txt(x,285,450,35,title,18,INK,True)

# Strategy: frozen means no parameter updates, not zero computation.
txt(68,327,460,26,'在指定 6B 预训练模型上，先练动作决策',13,MUTED)
rect(68,365,460,83,'E8F5F0',8)
txt(82,375,90,26,'更新参数',13,TEAL,True)
txt(178,375,338,61,'动作专家、状态 / 动作投影\n深度 / 视频查询、共享投影、对齐头',12.3,INK,leading=29)
rect(68,458,460,78,'F0F3F7',8)
txt(82,469,90,26,'冻结参数',13,MUTED,True)
txt(178,469,338,60,'视觉语言骨干（VLM）\n深度 / 视频教师；仍参与计算',12.5,INK,leading=28)
txt(68,549,460,25,'目标：根据图像、指令与状态，预测未来 50 步动作',12.2,INK)
txt(68,577,460,25,'保留当前 / 未来深度、视频辅助损失（各 0.004）',12.2,INK)
txt(68,608,460,25,'每卡 12 × 8 卡 × 梯度累计 1 = 有效 batch 96',12.5,INK)
txt(68,637,460,25,'基础 LR 5e-5 · warmup 2% · 余弦下降至 1e-5',12.2,INK)
txt(68,667,460,25,'梯度检查点开 · 混合精度开 · epoch 末保存至 OBS',12.0,TEAL)

# Time comparison: common linear scale.
txt(592,327,434,26,'同样 5717 步 / 轮；近期训练持续推进',13,MUTED)
max_w=270
for y,label,s,clr,win in [(370,'优化前',old,'AAB9C8','500 步'),(435,'当前',sec,TEAL,'100 步')]:
 txt(592,y,104,25,label,13,INK)
 txt(742,y,284,25,f'{s:.2f} 秒/步  ≈ {s*5717/3600:.1f} 小时/轮',12.4,clr if clr==TEAL else MUTED,align='right')
 rect(592,y+32,max_w*(s/old),15,clr,4)
 txt(878,y+27,148,22,f'{win}计时窗口',10.5,MUTED,align='right')
txt(592,494,434,43,f'每步耗时减少约 {saving:.0f}%',23,TEAL,True)
txt(592,536,434,23,'视频读取优化后实测提速；不同运行窗口对比',11.4,MUTED)
txt(592,572,434,29,f'当前 {step} / {total} 步  ·  {100*step/total:.1f}%',15,INK,True)
rect(592,607,434,10,LINE,3);rect(592,607,434*step/total,10,TEAL,3)
txt(592,630,434,24,f'动作 loss {latest["vla_loss"]:.3f} · 尚无任务成功率评测',12.1,MUTED)
txt(592,658,434,26,f'预计剩余约 {T["remaining_hours"]:.1f} 小时',14,INK)
txt(592,690,434,22,'不含启动、保存与上传；短期速度仍可能变化',10.7,MUTED)

# Resource readings: identical 0-100 percent scales on each card.
txt(1090,327,442,26,'8 × Ascend 910B1 · 每卡 64 GiB',13,MUTED)
txt(1090,366,210,52,f'{N["overall_aicore_mean"]:.1f}%',33,AMBER,True)
txt(1320,366,212,52,'96–97%',33,INK,True)
txt(1090,420,210,24,'AI Core 采样均值',12,MUTED)
txt(1320,420,212,24,'显存占用',12,MUTED)
txt(1090,459,442,23,'各卡 AI Core 均值（统一 0–100% 刻度）',11.4,MUTED)
for k,c in enumerate(N['cards']):
 col=k//4;row=k%4;x=1090+col*228;y=488+row*24
 txt(x,y,34,23,str(k)+' 卡',10.4,MUTED)
 rect(x+41,y+9,108,7,LINE,2);rect(x+41,y+9,108*c['aicore_mean']/100,7,AMBER,2)
 txt(x+152,y,57,23,f'{c["aicore_mean"]:.1f}%',10.4,INK,align='right')
txt(1090,597,442,30,'显存接近满 ≠ 计算持续跑满',15,AMBER,True)
lo=min(c['power_w_mean'] for c in N['cards']);hi=max(c['power_w_mean'] for c in N['cards'])
txt(1090,635,442,24,f'健康均为 OK · 50–57℃ · 约 {lo:.0f}–{hi:.0f} W / 卡',11.6,MUTED)
duration=(datetime.fromisoformat(N['sample_end'])-datetime.fromisoformat(N['sample_start'])).total_seconds()
txt(1090,664,442,23,f'{duration:.0f} 秒 / {N["sample_count"]} 次采样；读数不等于理论算力利用率',11.2,MUTED)
txt(1090,691,442,22,'低利用率的具体原因，仍需性能分析确认',10.8,MUTED)

rect(48,736,1504,114,INK,10)
txt(68,748,1464,28,'下一步讨论｜先定位等待点，再决定改什么',15,WHITE,True)
for x,title,body in [
 (68,'① 分段计时','分开看取数据、计算、通信，优先找等待。'),
 (560,'② 检查 NPU 算子','排查注意力 / MoE 热点，验证融合实现收益。'),
 (1052,'③ 控制试验范围','有显存余量再试少量减少重计算；不盲增 batch。')]:
 txt(x,785,472,26,title,13,'7ED6C1',True)
 txt(x,816,472,25,body,11.4,'D8E3ED')
txt(48,866,1504,23,'依据：实际启动配置、train.log / loss.jsonl、npu-smi 连续采样。当前为训练早期观察，最终效果需 randomized 仿真评测。',10.5,MUTED)

notes='''数据与口径：
- 本页是单次现场快照，详细原始采样保存在同目录 metrics_snapshot.json。
- TRAIN_EXPERT_ONLY=true；动作专家之外，状态/动作投影、深度/视频查询、共享投影与对齐头仍可训练。
- VLM 参数冻结不代表全程 no_grad；辅助教师冻结，但继续前向提供监督。
- 1 epoch，5717 updates；8 cards，microbatch 12，gradient_accumulation_steps 1，全局 batch 96。
- Muon，weight_decay 0.001；基础 LR 5e-5 到 1e-5，2% warmup=114 steps。MoE expert LR 另乘 sqrt(8)，并非所有参数都使用相同 LR。
- 梯度检查点与 mixed precision 均打开；不能称为纯 FP32 训练。
- 当前数据是朋友提供的 LeRobot v3 转换数据，2500 episodes、548893 frames；官方 clean 来源尚未独立确认。
- 每轮末保存 checkpoint，上传至 obs://yw-2030-extern/Partner_Qin/robot_1014/robotwin_1ep_fast/，校验后清理本地。
- 速度为日志时间戳间隔 / 更新数；新运行近100步、旧运行末500步，非严格控制变量基准。机器背景负载也可能影响结果。
- 视频优化限定解码线程并只转换被选中的帧；不修改原始数据，已验证抽样帧逐像素一致。
- 单轮/剩余时间按近期速度线性推算，未包含初始化、保存、上传和评测。
- npu-smi AI Core 读数不是 MFU，不应直接等同于理论 FLOPS 利用率。显存占用不等于计算忙碌程度。
- 下一步建议是待验证方向；未改变学习目标、未关闭辅助损失、未直接增加 batch 或梯度累计。
- 本次生成报告没有停止或修改训练。
证据：
logs/robotwin_1ep_fast/lingbotvla_cli.yaml
logs/robotwin_1ep_fast/preparation.json
logs/robotwin_1ep_fast/previous_run_baseline.json
logs/robotwin_1ep_fast/train.log
logs/robotwin_1ep_fast/loss.jsonl
'''
slide.notes_slide.notes_text_frame.text=notes
out=ROOT/'LingBot训练分析_1epoch_20261016.pptx'
prs.save(out)
im.save(ROOT/'slide_preview.png')
# Structural and text-fit checks. All visible information stays editable in PPTX.
check=Presentation(out)
assert len(check.slides)==1
assert all(x>=0 and y>=0 and x+w<=W+1 and y+h<=H+1 for _,x,y,w,h,*_ in objects)
assert '128' in '\n'.join(s.text for s in check.slides[0].shapes if s.has_text_frame) or str(step) in '\n'.join(s.text for s in check.slides[0].shapes if s.has_text_frame)
(ROOT/'slide_validation.json').write_text(json.dumps({'slides':len(check.slides),'editable_shapes':len(check.slides[0].shapes),'text_fit_warnings':warnings,'snapshot_at':D['captured_at'],'file_bytes':out.stat().st_size},ensure_ascii=False,indent=2))
print(json.dumps({'pptx':str(out),'preview':str(ROOT/'slide_preview.png'),'warnings':warnings},ensure_ascii=False))
